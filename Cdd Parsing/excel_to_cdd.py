"""Write the edits of an Excel file exported by cdd_parser.py back into a CDD file.

The original CDD is used as the template: every row of the Excel file points to its element in the CDD
through the ref columns ("CDD Ref", "DID Ref", "Element Ref"), and only the values you changed are
rewritten. Everything else in the CDD, including its formatting, stays byte for byte the same, so the
result opens in CANdela Studio like the original.

What can be edited (every other column is read-only):
    Service  Sessions, Security, Addressing
    DID      DID, Name, Qualifier, Read/Write Sessions and Security,
             Data Element, Element Qualifier, Conversion, Unit, Value Table texts
    DTC      DTC (or DTC Code), Description and every DTC record field the DTC already has
    RID      RID, Name, Qualifier, Description, Sessions, Security

Every problem is listed with its sheet, row and column. If there is any error, no CDD is written.

Usage:
    python excel_to_cdd.py <original.cdd> <edited.xlsx> [-o <new.cdd>] [--report <report.txt>]
"""

import argparse
import os
import re
import sys
import xml.parsers.expat
from pathlib import Path

from cdd_parser import (
    DID_REF_COLUMN, ELEMENT_REF_COLUMN, READ_DID_SID, REF_COLUMN, ROUTINE_CONTROL_SID, WRITE_DID_SID,
    CddFile, build_export, excel_sheets, nested_structure, structure_elements, text,
)

SHEETS = ("Service", "DID", "DTC", "RID")
HEX_COLUMNS = {"DID": 2, "DTC": 3, "RID": 2, "SID": 1}
DID_LEVEL_COLUMNS = ("DID", "Name", "Qualifier", "Variants", "Length (bytes)", "Read (22)", "Read Sessions",
                     "Read Security", "Write (2E)", "Write Sessions", "Write Security")
DID_DATA_SIDS = {0x22, 0x2E, 0x2F}  # services whose identifier is the DID number
QUALIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
EDITABLE = {
    "Service": {"Sessions", "Security", "Addressing"},
    "DID": {"DID", "Name", "Qualifier", "Read Sessions", "Read Security", "Write Sessions", "Write Security",
            "Data Element", "Element Qualifier", "Conversion", "Unit", "Value Table"},
    "DTC": {"DTC", "DTC Code", "Description"},  # plus every DTC record field
    "RID": {"RID", "Name", "Qualifier", "Description", "Sessions", "Security"},
}


class Report:
    def __init__(self):
        self.errors, self.warnings, self.changes = [], [], []

    def error(self, where, message):
        self.errors.append(f"ERROR    {where}: {message}")

    def warning(self, where, message):
        self.warnings.append(f"WARNING  {where}: {message}")

    def change(self, where, message):
        self.changes.append(f"CHANGED  {where}: {message}")

    def lines(self):
        summary = f"{len(self.errors)} error(s), {len(self.warnings)} warning(s), {len(self.changes)} change(s)"
        return self.errors + self.warnings + self.changes + [summary]


# ---------- reading and comparing rows ----------

def norm(value):
    """Cell value as stripped text, the way the exporter wrote it."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").strip()


def short(value, limit=70):
    """One-line preview of a cell value for the report."""
    value = " / ".join(line.strip() for line in str(value).split("\n") if line.strip())
    return value if len(value) <= limit else value[:limit - 3] + "..."


def parse_hex(value, n_bytes):
    """'F190', '0xF190' or '0x f1 90' -> int; raises ValueError with a readable message."""
    digits = value.lower().replace("0x", "").replace(" ", "")
    if not re.fullmatch(r"[0-9a-f]+", digits):
        raise ValueError(f"'{value}' is not a hex number")
    number = int(digits, 16)
    if number >= 1 << (8 * n_bytes):
        raise ValueError(f"'{value}' does not fit in {n_bytes} byte(s)")
    return number


def same(column, a, b):
    if a == b:
        return True
    if column in HEX_COLUMNS:  # Excel may turn '0460' into 460
        try:
            return parse_hex(a, 8) == parse_hex(b, 8)
        except ValueError:
            return False
    return False


def read_workbook(path, report):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Reading Excel needs openpyxl: pip install openpyxl")
    workbook = load_workbook(path, data_only=True)
    sheets = {}
    for name in SHEETS:
        if name not in workbook.sheetnames:
            report.error(f"{name} sheet", "the sheet is missing")
            continue
        rows = list(workbook[name].iter_rows(values_only=True))
        header = [norm(h) for h in rows[0]] if rows else []
        records = []
        for number, values in enumerate(rows[1:], start=2):
            if all(norm(v) == "" for v in values):
                continue
            records.append((number, {h: norm(v) for h, v in zip(header, values) if h}))
        sheets[name] = (header, records)
    return sheets


def row_key(sheet, row):
    if sheet == "DID":
        return row.get(DID_REF_COLUMN, ""), row.get(ELEMENT_REF_COLUMN, "")
    return row.get(REF_COLUMN, "")


def describe(sheet, number, row):
    """'DTC sheet row 5 (DTC C41881)' for messages."""
    label = {
        "Service": lambda r: f"{r.get('Variant')} / {r.get('Diag Instance')} / {r.get('Service')}",
        "DID": lambda r: f"DID {r.get('DID')}" + (f" / {r.get('Data Element')}" if r.get("Data Element") else ""),
        "DTC": lambda r: f"DTC {r.get('DTC')}",
        "RID": lambda r: f"{r.get('Variant')} / RID {r.get('RID')}",
    }[sheet](row)
    return f"{sheet} sheet row {number} ({label})"


def match_rows(sheet, original_rows, header, edited_rows, report):
    """Pair each edited row with its original row by ref. Returns [(row number, original, edited)]."""
    ref_columns = (DID_REF_COLUMN, ELEMENT_REF_COLUMN) if sheet == "DID" else (REF_COLUMN,)
    missing = [c for c in ref_columns if c not in header]
    if missing:
        report.error(f"{sheet} sheet", f"column {', '.join(missing)} is missing. Export the Excel file again with "
                                       "the current cdd_parser.py and make your edits in that file")
        return []
    originals = {row_key(sheet, row): row for row in original_rows}
    seen, pairs = {}, []
    for number, row in edited_rows:
        key = row_key(sheet, row)
        where = describe(sheet, number, row)
        if not all(key if isinstance(key, tuple) else (key,)):
            report.error(where, "new rows are not supported; only existing rows can be edited")
        elif key not in originals:
            report.error(where, f"unknown ref {key}; the row does not belong to this CDD")
        elif key in seen:
            report.error(where, f"same ref as row {seen[key]}; rows must not be copied")
        else:
            seen[key] = number
            pairs.append((number, originals[key], row))
    for key, row in originals.items():
        if key not in seen:
            report.error(f"{sheet} sheet ({describe(sheet, '-', row).split('(', 1)[1][:-1]})",
                         "the row was deleted; deleting is not supported, the row must stay in the file")
    return pairs


def changed_cells(original, edited):
    """Columns whose value in the edited row differs from the original row."""
    return {column: value for column, value in edited.items()
            if not same(column, original.get(column, ""), value)}


# ---------- planned edits on the CDD ----------

class Edits:
    """Attribute and text changes on CDD elements, with conflict detection."""

    def __init__(self, report):
        self.report = report
        self.attrs = {}   # element -> {attribute: (value, where)}
        self.texts = {}   # element -> (text, where)

    def set_attr(self, elem, attr, value, where):
        value = str(value)
        current = self.attrs.get(elem, {}).get(attr)
        if current and current[0] != value:
            self.report.error(where, f"conflicts with {current[1]}, which sets the same value to '{current[0]}'")
        elif elem.get(attr) != value:
            self.attrs.setdefault(elem, {})[attr] = (value, where)

    def set_text(self, elem, value, where):
        current = self.texts.get(elem)
        if current and current[0] != value:
            self.report.error(where, f"conflicts with {current[1]}, which sets the same text to '{current[0]}'")
        elif (elem.text or "").strip() != value:
            self.texts[elem] = (value, where)


def text_node(elem, tag="NAME"):
    """The element holding the text of NAME/DESC/TEXT (its first TUV), or None."""
    node = elem.find(tag) if elem is not None else None
    if node is None:
        return None
    tuv = node.find("TUV")
    return tuv if tuv is not None else node


class Applier:
    def __init__(self, cdd, report):
        self.cdd, self.report = cdd, report
        self.edits = Edits(report)
        # DIDs get their own index: CANdela reuses DID ids for diagnostic instances.
        self.dids = {did.get("id"): did for did in cdd.doc.find("DIDS").findall("DID")}
        self.state_index = {}  # (group spec, state name) -> 1-based mayBeExec index
        for index, (spec, name) in enumerate(cdd.states, start=1):
            self.state_index[(spec, name)] = index

    # ----- shared helpers -----

    def set_name(self, elem, value, where, tag="NAME"):
        node = text_node(elem, tag)
        if node is None:
            self.report.error(where, f"the element has no {tag} to change")
        elif not value:
            self.report.error(where, "the text must not be empty")
        else:
            self.edits.set_text(node, value, where)

    def set_qualifier(self, elem, value, where):
        node = elem.find("QUAL")
        if not QUALIFIER.match(value):
            self.report.error(where, f"'{value}' is not a valid qualifier (letters, digits and _, not starting with "
                                     "a digit)")
        elif node is None:
            self.report.error(where, "the element has no qualifier to change")
        else:
            self.edits.set_text(node, value, where)

    def may_be_exec(self, sessions, security, where):
        """Build a mayBeExec value such as '(1,3,4,6)' from the Sessions and Security cells."""
        indexes = []
        for spec, cell in (("session", sessions), ("security", security)):
            valid = [name for (s, name) in self.state_index if s == spec]
            if cell == "All":
                indexes += [self.state_index[(spec, name)] for name in valid]
                continue
            if cell in ("", "-"):
                continue
            for name in (n.strip() for n in cell.split(",")):
                if (spec, name) not in self.state_index:
                    self.report.error(where, f"unknown {spec} state '{name}'; valid: All, -, {', '.join(valid)}")
                    return None
                indexes.append(self.state_index[(spec, name)])
        return "(" + ",".join(str(i) for i in sorted(set(indexes))) + ")"

    def set_access(self, services, sessions, security, where):
        value = self.may_be_exec(sessions, security, where)
        if value is not None:
            for service in services:
                self.edits.set_attr(service, "mayBeExec", value, where)

    # ----- sheets -----

    def service(self, where, original, changes):
        service = self.cdd.ids.get(original[REF_COLUMN].split("@")[0])
        if "Sessions" in changes or "Security" in changes:
            self.set_access([service], changes.get("Sessions", original.get("Sessions", "")),
                            changes.get("Security", original.get("Security", "")), f"{where}, column Sessions/Security")
        if "Addressing" in changes:
            cell = f"{where}, column Addressing"
            modes = {m.strip().lower() for m in changes["Addressing"].split(",") if m.strip()}
            if not modes or modes - {"physical", "functional"}:
                self.report.error(cell, "use Physical, Functional or both, separated by a comma")
            else:
                self.edits.set_attr(service, "phys", "1" if "physical" in modes else "0", cell)
                self.edits.set_attr(service, "func", "1" if "functional" in modes else "0", cell)

    def did_services(self, did, number):
        """Diagnostic instances that use this DID number, in the variants that list the DID."""
        found = []
        for variant in self.cdd.variants():
            if not any(ref.get("didRef") == did.get("id") for ref in variant.iter("DIDREF")):
                continue
            for _, inst in self.cdd.diag_instances(variant):
                sv = inst.find("STATICVALUE")
                if sv is None or int(sv.get("v")) != number or inst in (i for i, _ in found):
                    continue
                sids = {s: self.cdd.request_bytes(inst, s)[0] for s in inst.findall("SERVICE")}
                if DID_DATA_SIDS & set(sids.values()):
                    found.append((inst, sids))
        return found

    def did(self, where, did, original, changes):
        number = int(did.get("n"))
        users = self.did_services(did, number)
        if "DID" in changes:
            cell = f"{where}, column DID"
            try:
                new = parse_hex(changes["DID"], 2)
            except ValueError as exc:
                self.report.error(cell, str(exc))
            else:
                self.edits.set_attr(did, "n", new, cell)
                for inst, _ in users:
                    self.edits.set_attr(inst.find("STATICVALUE"), "v", new, cell)
        if "Name" in changes:
            self.set_name(did, changes["Name"], f"{where}, column Name")
        if "Qualifier" in changes:
            self.set_qualifier(did, changes["Qualifier"], f"{where}, column Qualifier")
        for label, sid in (("Read", READ_DID_SID), ("Write", WRITE_DID_SID)):
            if f"{label} Sessions" in changes or f"{label} Security" in changes:
                cell = f"{where}, column {label} Sessions/Security"
                services = [s for inst, sids in users for s, s_sid in sids.items() if s_sid == sid]
                if not services:
                    self.report.error(cell, f"the DID has no {label.lower()} service, so it has no access to change")
                    continue
                self.set_access(services, changes.get(f"{label} Sessions", original.get(f"{label} Sessions", "")),
                                changes.get(f"{label} Security", original.get(f"{label} Security", "")), cell)

    def element(self, where, ref, original, changes):
        obj, dtype = self.resolve_element(ref)
        if obj is None:
            self.report.error(where, f"element ref {ref} not found in the CDD")
            return
        if "Data Element" in changes:
            name = changes["Data Element"].split(" / ")[-1].strip()
            self.set_name(obj, name, f"{where}, column Data Element")
        if "Element Qualifier" in changes:
            self.set_qualifier(obj, changes["Element Qualifier"], f"{where}, column Element Qualifier")
        if dtype is None:
            return
        if "Unit" in changes:
            cell = f"{where}, column Unit"
            unit = dtype.find("PVALUETYPE/UNIT")
            if unit is None:
                self.report.error(cell, "the data type has no unit; adding a unit is not supported")
            else:
                self.edits.set_text(unit, changes["Unit"], cell)
        if "Conversion" in changes:
            self.conversion(dtype, changes["Conversion"], f"{where}, column Conversion")
        if "Value Table" in changes:
            self.value_table(dtype, changes["Value Table"], f"{where}, column Value Table")

    def resolve_element(self, ref):
        did_id, _, path = ref.partition(":")
        did = self.dids.get(did_id)
        structure = did.find("STRUCTURE") if did is not None else None
        obj = dtype = None
        try:
            for index in path.split("."):
                obj = structure_elements(structure)[int(index)]
                dtype = None if obj.tag == "GAPDATAOBJ" else self.cdd.ids.get(obj.get("dtref"))
                structure = nested_structure(obj, dtype)
        except (ValueError, IndexError):
            return None, None
        return obj, dtype

    def conversion(self, dtype, cell_text, where):
        comps = dtype.findall("COMP")
        lines = [line.strip() for line in cell_text.split("\n") if line.strip()]
        if len(lines) != len(comps):
            self.report.error(where, f"the data type has {len(comps)} conversion line(s); adding or removing "
                                     "conversions is not supported")
            return
        pattern = re.compile(r"^(?:raw\s+(?P<s>-?[\d.]*)\.\.(?P<e>-?[\d.]*)\s*:\s*)?(?P<rest>.*)$")
        for comp, line in zip(comps, lines):
            match = pattern.match(line)
            values = {}
            for part in match.group("rest").split(","):
                key, _, value = part.strip().partition(" ")
                values[key] = value.strip()
            if set(values) - {"factor", "divisor", "offset"} or not values.get("factor") or not values.get("offset"):
                self.report.error(where, f"'{line}' should look like 'raw 0..255: factor 0.5, divisor 1, offset 0'")
                continue
            try:
                for value in values.values():
                    float(value)
                for bound in (match.group("s"), match.group("e")):
                    if bound:
                        float(bound)
            except ValueError:
                self.report.error(where, f"'{line}' contains a value that is not a number")
                continue
            if values.get("divisor") is None and comp.get("div") is not None:
                values["divisor"] = "1"
            for attr, key in (("f", "factor"), ("div", "divisor"), ("o", "offset")):
                if key in values:
                    self.edits.set_attr(comp, attr, values[key], where)
            for attr in ("s", "e"):
                bound = match.group(attr)
                if bound is not None and (bound or comp.get(attr) is not None):
                    if not bound:
                        self.report.error(where, f"removing the raw range limit in '{line}' is not supported")
                    else:
                        self.edits.set_attr(comp, attr, bound, where)

    def value_table(self, dtype, cell_text, where):
        maps = dtype.findall("TEXTMAP")
        lines = [line for line in cell_text.split("\n") if line.strip()]
        if len(lines) != len(maps):
            self.report.error(where, f"the value table has {len(maps)} entries; adding or removing entries is not "
                                     "supported")
            return
        for mapping, line in zip(maps, lines):
            raw, sep, label = line.partition("=")
            expected = mapping.get("s") if mapping.get("s") == mapping.get("e") else \
                f"{mapping.get('s')}..{mapping.get('e')}"
            if not sep or raw.strip() != expected:
                self.report.error(where, f"'{line}' must keep its raw value '{expected} = ...'; only the text can "
                                         "change")
                continue
            self.set_name(mapping, label.strip(), where, tag="TEXT")

    def dtc(self, where, record, original, changes):
        number = None
        if "DTC" in changes:
            try:
                number = parse_hex(changes["DTC"], 3)
            except ValueError as exc:
                self.report.error(f"{where}, column DTC", str(exc))
        elif "DTC Code" in changes:
            number = parse_dtc_code(changes["DTC Code"])
            if number is None:
                self.report.error(f"{where}, column DTC Code", f"'{changes['DTC Code']}' is not a DTC code like "
                                                               "U0418-81")
        if number is not None:
            self.edits.set_attr(record, "v", number, f"{where}, column DTC")
        if "Description" in changes:
            self.set_name(record, changes["Description"], f"{where}, column Description", tag="TEXT")
        items = {}
        for item in record:
            template = self.cdd.ref(item, "idref")
            if template is not None:
                items[text(template)] = (item, template)
        for column, value in changes.items():
            if column in ("DTC", "DTC Code", "Description") or column not in items and column not in self.dtc_fields():
                continue
            cell = f"{where}, column {column}"
            if column not in items:
                self.report.error(cell, "the DTC has no such field yet; adding a field is not supported")
                continue
            item, template = items[column]
            if not value:
                self.report.error(cell, "removing a field is not supported")
            elif item.tag == "TRRECORDITEM":
                self.set_name(item, value, cell, tag="TEXT")
            elif item.tag == "ENUMRECORDITEM":
                labels = {text(e, None): e.get("v") for e in template.findall("ETAG")}
                if value in labels:
                    self.edits.set_attr(item, "v", labels[value], cell)
                else:
                    self.report.error(cell, f"'{value}' is not one of: {', '.join(labels)}")
            elif value.isdigit():
                self.edits.set_attr(item, "v", value, cell)
            else:
                self.report.error(cell, f"'{value}' is not a whole number")

    def dtc_fields(self):
        return {text(item) for tmpl in self.cdd.doc.iter("RECORDTMPL") for item in tmpl if item.get("id")}

    def rid(self, where, inst, original, changes):
        if "RID" in changes:
            cell = f"{where}, column RID"
            try:
                self.edits.set_attr(inst.find("STATICVALUE"), "v", parse_hex(changes["RID"], 2), cell)
            except ValueError as exc:
                self.report.error(cell, str(exc))
        if "Name" in changes:
            self.set_name(inst, changes["Name"], f"{where}, column Name")
        if "Qualifier" in changes:
            self.set_qualifier(inst, changes["Qualifier"], f"{where}, column Qualifier")
        if "Description" in changes:
            cell = f"{where}, column Description"
            if text_node(inst, "DESC") is None:
                self.report.error(cell, "the routine has no description; adding one is not supported")
            else:
                self.set_name(inst, changes["Description"], cell, tag="DESC")
        if "Sessions" in changes or "Security" in changes:
            services = [s for s in inst.findall("SERVICE")
                        if self.cdd.request_bytes(inst, s)[0] == ROUTINE_CONTROL_SID]
            self.set_access(services, changes.get("Sessions", original.get("Sessions", "")),
                            changes.get("Security", original.get("Security", "")), f"{where}, column Sessions/Security")


def parse_dtc_code(code):
    """'U0418-81' -> 0xC41881, or None."""
    match = re.fullmatch(r"([PCBU])([0-3])([0-9A-F])([0-9A-F]{2})-?([0-9A-F]{2})", code.strip().upper())
    if not match:
        return None
    letter, d1, d2, middle, ftb = match.groups()
    high = ("PCBU".index(letter) << 6) | (int(d1) << 4) | int(d2, 16)
    return (high << 16) | (int(middle, 16) << 8) | int(ftb, 16)


# ---------- writing the CDD ----------

def element_positions(data):
    """Byte offsets of each element's start tag and end tag, in document order."""
    parser = xml.parsers.expat.ParserCreate()
    starts, ends, stack = [], [], []

    def on_start(name, attrs):
        stack.append(len(starts))
        starts.append(parser.CurrentByteIndex)
        ends.append(None)

    def on_end(name):
        ends[stack.pop()] = parser.CurrentByteIndex

    parser.StartElementHandler, parser.EndElementHandler = on_start, on_end
    parser.Parse(data, True)
    return starts, ends


def start_tag_end(data, start):
    """Index of the '>' that closes the start tag beginning at `start` (quotes respected)."""
    quote = None
    for index in range(start, len(data)):
        char = data[index:index + 1]
        if quote:
            if char == quote:
                quote = None
        elif char in (b"'", b'"'):
            quote = char
        elif char == b">":
            return index
    raise ValueError("unterminated start tag")


def escape_text(value, newline):
    # Same entity style as CANdela: &#38; &#60; &#62;
    value = value.replace("&", "&#38;").replace("<", "&#60;").replace(">", "&#62;")
    return value.replace("\n", newline)


def escape_attr(value, quote):
    value = value.replace("&", "&#38;").replace("<", "&#60;")
    return value.replace(quote, "&#39;" if quote == "'" else "&#34;")


def patch_start_tag(tag, changes):
    """Set attributes inside a raw start tag, keeping its quote style."""
    for attr, value in changes.items():
        pattern = re.compile(r"(\s" + re.escape(attr) + r"\s*=\s*)(['\"])(.*?)\2", re.S)
        match = pattern.search(tag)
        if match:
            quote = match.group(2)
            tag = tag[:match.start()] + match.group(1) + quote + escape_attr(value, quote) + quote + tag[match.end():]
        else:
            quote = "'" if "='" in tag else '"'
            end = len(tag) - (2 if tag.endswith("/>") else 1)
            tag = tag[:end] + f" {attr}={quote}{escape_attr(value, quote)}{quote}" + tag[end:]
    return tag


def write_cdd(data, root, edits, path):
    starts, ends = element_positions(data)
    position = {elem: index for index, elem in enumerate(root.iter())}
    replacements = []
    for elem, changes in edits.attrs.items():
        index = position[elem]
        end = start_tag_end(data, starts[index])
        tag = data[starts[index]:end + 1].decode("utf-8")
        new = patch_start_tag(tag, {attr: value for attr, (value, _) in changes.items()})
        replacements.append((starts[index], end + 1, new.encode("utf-8")))
    for elem, (value, _) in edits.texts.items():
        index = position[elem]
        open_end = start_tag_end(data, starts[index])
        if data[open_end - 1:open_end] == b"/":  # <TAG/> becomes <TAG>text</TAG>
            tag = data[starts[index]:open_end - 1].rstrip()
            name = elem.tag.encode("utf-8")
            new = tag + b">" + escape_text(value, "\n").encode("utf-8") + b"</" + name + b">"
            replacements.append((starts[index], open_end + 1, new))
            continue
        old = data[open_end + 1:ends[index]]
        newline = "\r\n" if b"\r\n" in old else "\n"
        replacements.append((open_end + 1, ends[index], escape_text(value, newline).encode("utf-8")))
    replacements.sort()
    for (_, end_a, _), (start_b, _, _) in zip(replacements, replacements[1:]):
        if start_b < end_a:
            raise RuntimeError("overlapping edits")
    out, last = [], 0
    for start, end, new in replacements:
        out += [data[last:start], new]
        last = end
    out.append(data[last:])
    Path(path).write_bytes(b"".join(out))


# ---------- main flow ----------

def logical_cells(sheet, rows):
    """{(ref, column): value}; DID-level columns are keyed by the DID ref so each DID counts once."""
    cells = {}
    for number, row in rows:
        for column, value in row.items():
            if column in (REF_COLUMN, DID_REF_COLUMN, ELEMENT_REF_COLUMN):
                continue
            if sheet == "DID" and column in DID_LEVEL_COLUMNS:
                key = (row.get(DID_REF_COLUMN), column)
            else:
                key = (row_key(sheet, row), column)
            cells.setdefault(key, []).append((number, row, value))
    return cells


def convert(cdd_path, excel_path, output_path, report):
    original_tree, _ = build_export(cdd_path)
    original = {name: [(0, {c: norm(v) for c, v in row.items()}) for row in rows]
                for name, rows in excel_sheets(original_tree).items()}
    edited = read_workbook(excel_path, report)
    cdd = CddFile(cdd_path)
    applier = Applier(cdd, report)

    pairs = {}
    for sheet in SHEETS:
        if sheet in edited:
            header, rows = edited[sheet]
            pairs[sheet] = match_rows(sheet, [r for _, r in original[sheet]], header, rows, report)

    # Columns computed from the CDD (lengths, types, request bytes, ...) cannot be edited.
    dtc_fields = applier.dtc_fields()
    for sheet, sheet_pairs in pairs.items():
        for number, orig, row in sheet_pairs:
            for column, value in changed_cells(orig, row).items():
                if column not in EDITABLE[sheet] and not (sheet == "DTC" and column in dtc_fields):
                    report.error(f"{describe(sheet, number, row)}, column {column}",
                                 f"this column is read-only (it is computed from the CDD); found '{short(value)}', "
                                 f"expected '{short(orig.get(column, ''))}'")

    # DID-level columns repeat on every element row: one changed value per DID, or an error.
    did_changes = {}
    for number, orig, row in pairs.get("DID", []):
        for column in DID_LEVEL_COLUMNS:
            if column in row and not same(column, orig.get(column, ""), row[column]):
                did_changes.setdefault(orig[DID_REF_COLUMN], {}).setdefault(column, {})[row[column]] = number
    for number, orig, row in pairs.get("DID", []):
        changes = changed_cells(orig, row)
        where = describe("DID", number, row)
        for column in DID_LEVEL_COLUMNS:
            changes.pop(column, None)
        if orig.get(ELEMENT_REF_COLUMN):
            applier.element(where, orig[ELEMENT_REF_COLUMN], orig, changes)
    for did_ref, columns in did_changes.items():
        did = applier.dids[did_ref]
        rows = {n: r for n, o, r in pairs["DID"] if o[DID_REF_COLUMN] == did_ref}
        first = min(rows)
        where = describe("DID", first, rows[first])
        changes = {}
        for column, values in columns.items():
            if len(values) > 1:
                report.error(where, f"column {column} has different new values in rows {', '.join(map(str, sorted(values.values())))}")
            else:
                changes[column] = next(iter(values))
        orig = next(o for n, o, r in pairs["DID"] if o[DID_REF_COLUMN] == did_ref)
        applier.did(where, did, orig, changes)

    for sheet in ("Service", "DTC", "RID"):
        for number, orig, row in pairs.get(sheet, []):
            changes = changed_cells(orig, row)
            if not changes:
                continue
            where = describe(sheet, number, row)
            # Services and routines are shared between variants, so their ref is '<element id>@<variant id>'.
            elem = cdd.ids.get(orig[REF_COLUMN].split("@")[0])
            if sheet == "Service":
                applier.service(where, orig, changes)
            elif sheet == "DTC":
                applier.dtc(where, elem, orig, changes)
            else:
                applier.rid(where, elem, orig, changes)

    check_duplicates(cdd, applier.dids, applier.edits, report)
    if report.errors:
        return False

    # Write to a temporary file, export it again and check every changed cell landed.
    data = Path(cdd_path).read_bytes()
    temp = f"{output_path}.tmp"
    write_cdd(data, cdd.root, applier.edits, temp)
    try:
        new_tree, _ = build_export(temp)
    except Exception as exc:  # the written file must still be valid XML
        os.remove(temp)
        report.error("output", f"the written CDD could not be read back: {exc}")
        return False
    new = {name: [(0, {c: norm(v) for c, v in row.items()}) for row in rows]
           for name, rows in excel_sheets(new_tree).items()}
    verify(original, edited, new, report)
    if report.errors:
        os.remove(temp)
        return False
    os.replace(temp, output_path)
    return True


def verify(original, edited, new, report):
    """Every changed cell must read back from the new CDD; other cells that changed are reported as side effects."""
    def unit(key):  # the DID, DTC, RID or service a cell belongs to
        return key[0][0] if isinstance(key[0], tuple) else key[0]

    def first_value(cells, key):
        return cells.get(key, [(0, {}, "")])[0][2]

    for sheet in SHEETS:
        if sheet not in edited:
            continue
        before = logical_cells(sheet, original[sheet])
        after = logical_cells(sheet, new[sheet])
        wanted = logical_cells(sheet, edited[sheet][1])
        edited_units = set()
        for key, entries in wanted.items():
            old, now = first_value(before, key), first_value(after, key)
            column = key[1]
            for number, row, value in entries:
                if same(column, old, value):
                    continue
                edited_units.add(unit(key))
                where = f"{describe(sheet, number, row)}, column {column}"
                if same(column, now, value):
                    report.change(where, f"'{short(old)}' -> '{short(value)}'")
                else:
                    report.error(where, "this column cannot be changed (it is computed from the CDD or the change "
                                        f"is not supported); the CDD still gives '{short(now)}'")
                break
        for key, entries in before.items():
            old, now = first_value(before, key), first_value(after, key)
            if old != now and unit(key) not in edited_units:
                _, row, _ = entries[0]
                label = describe(sheet, "-", row).split("(", 1)[1][:-1]
                report.warning(f"{sheet} sheet ({label}), column {key[1]}",
                               f"also changed from '{short(old)}' to '{short(now)}' by another edit (shared data type, DID or "
                               "RID number, or similar)")


def check_duplicates(cdd, dids_by_id, edits, report):
    """New DID numbers must stay unique per variant, DTC numbers unique per fault memory."""
    def value(elem, attr):
        return edits.attrs.get(elem, {}).get(attr, (elem.get(attr), None))[0]

    for variant in cdd.variants():
        numbers = {}
        for ref in variant.iter("DIDREF"):
            did = dids_by_id.get(ref.get("didRef"))
            if did is not None:
                numbers.setdefault(int(value(did, "n")), []).append(did)
        for number, dids in numbers.items():
            if len(dids) > 1 and any(d in edits.attrs for d in dids):
                report.error(f"DID sheet (DID {number:04X})", f"variant {text(variant)} would have this DID "
                                                               f"{len(dids)} times")
    for record_dt in cdd.root.iter("RECORDDT"):
        numbers = {}
        for record in record_dt.findall("RECORD"):
            numbers.setdefault(int(value(record, "v")), []).append(record)
        for number, records in numbers.items():
            if len(records) > 1 and any(r in edits.attrs for r in records):
                report.error(f"DTC sheet (DTC {number:06X})", "the fault memory would have this DTC "
                                                               f"{len(records)} times")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Apply an edited Excel export to a CDD file.")
    parser.add_argument("cdd", help="the original .cdd file the Excel file was exported from")
    parser.add_argument("excel", help="the edited .xlsx file")
    parser.add_argument("-o", "--output", help="new .cdd file (default: <cdd name>_edited.cdd)")
    parser.add_argument("--report", help="also write the report to this text file")
    args = parser.parse_args(argv)
    output = args.output or str(Path(args.cdd).with_name(Path(args.cdd).stem + "_edited.cdd"))
    if Path(output).resolve() == Path(args.cdd).resolve():
        sys.exit("The output must not overwrite the original CDD")

    report = Report()
    ok = convert(args.cdd, args.excel, output, report)
    lines = report.lines() + ([f"Wrote {output}"] if ok else ["No CDD written because of the errors above"])
    print("\n".join(lines))
    if args.report:
        Path(args.report).write_text("\n".join(lines) + "\n", encoding="utf-8")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
