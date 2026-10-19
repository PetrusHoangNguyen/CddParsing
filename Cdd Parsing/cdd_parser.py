"""Parse a Vector CANdela Studio CDD file and export its diagnostic data to XML or Excel.

The output has four sections (XML) or sheets (Excel):
    Services  every diagnostic service of every variant, with request bytes and access rights
    DIDs      every data identifier with its data elements, data types and read/write access
    DTCs      every DTC of every fault memory, with the DTC record fields
    RIDs      every routine identifier with its sub-functions and parameters

Usage:
    python cdd_parser.py <file.cdd> [--format xml|excel] [-o <output file>]
"""

import argparse
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

ROUTINE_CONTROL_SID = 0x31
READ_DID_SID = 0x22
WRITE_DID_SID = 0x2E


def text(elem, tag="NAME"):
    """Return the (first language) text of a CANdela text element such as NAME or DESC."""
    if elem is None:
        return ""
    node = elem.find(tag) if tag else elem
    if node is None:
        return ""
    tuv = node.find("TUV")
    value = tuv.text if tuv is not None else node.text
    return (value or "").strip()


DATA_TYPE_KINDS = {
    "IDENT": "Identical",
    "LINCOMP": "Linear",
    "COMPTBL": "Piecewise linear",
    "TEXTTBL": "Text table",
    "STRUCTDT": "Structure",
    "EOSITERDT": "Iteration until end",
    "STATICITERRDT": "Static iteration",
    "MUXDT": "Multiplexer",
}
ENCODINGS = {
    "uns": "Unsigned", "sgn": "Signed", "asc": "ASCII", "bcd": "BCD",
    "utf": "Unicode", "dbl": "Float64", "flt": "Float32",
}
DISPLAY_FORMATS = {"hex": "Hex", "dec": "Decimal", "bin": "Binary", "text": "Text", "flt": "Float"}
BYTE_ORDERS = {"21": "Big endian (Motorola)", "12": "Little endian (Intel)"}


def value_range(elem):
    start, end = elem.get("s"), elem.get("e")
    return start if start == end else f"{start}..{end}"


def to_hex(value, n_bytes):
    return f"{value:0{n_bytes * 2}X}"


def dtc_code(number):
    """Convert a 3-byte UDS DTC number into its SAE J2012 form, e.g. 0xC41881 -> U0418-81."""
    high = (number >> 16) & 0xFF
    letter = "PCBU"[high >> 6]
    return f"{letter}{(high >> 4) & 0x3}{high & 0xF:X}{(number >> 8) & 0xFF:02X}-{number & 0xFF:02X}"


class CddFile:
    """In-memory view of a CDD file with an index of every element that has an id."""

    def __init__(self, path):
        self.path = Path(path)
        self.tree = ET.parse(self.path)
        self.root = self.tree.getroot()
        self.doc = self.root.find("ECUDOC")
        self.ids = {e.get("id"): e for e in self.root.iter() if e.get("id")}
        self.parent = {child: parent for parent in self.root.iter() for child in parent}
        # mayBeExec="(1,3,4,6)" counts the states of all state groups in document order.
        self.states = []
        for group in self.doc.iter("STATEGROUP"):
            for state in group.findall("STATE"):
                self.states.append((group.get("spec"), text(state)))

    # ---------- generic helpers ----------

    def ref(self, elem, attr):
        return self.ids.get(elem.get(attr)) if elem is not None else None

    def bit_length(self, dtref):
        """Return (min_bits, max_bits) of a data type, or (None, None) if unknown."""
        dtype = self.ids.get(dtref)
        cvt = dtype.find("CVALUETYPE") if dtype is not None else None
        if cvt is None:
            return None, None
        bl = int(cvt.get("bl", 0))
        if cvt.get("qty") == "field":
            return bl * int(cvt.get("minsz", 1)), bl * int(cvt.get("maxsz", 1))
        return bl, bl

    def access(self, service):
        """Return (sessions, security levels) in which a service may be executed."""
        may_be_exec = service.get("mayBeExec")
        template = self.ref(service, "tmplref")
        if may_be_exec is None and template is not None:
            may_be_exec = template.get("mayBeExec")
        if may_be_exec is None:
            return "All", "All"
        indexes = [int(i) for i in may_be_exec.strip("()").split(",") if i]
        picked = [self.states[i - 1] for i in indexes if 0 < i <= len(self.states)]
        sessions = ", ".join(name for spec, name in picked if spec == "session")
        security = ", ".join(name for spec, name in picked if spec == "security")
        return sessions or "-", security or "-"

    def protocol_service(self, service):
        """SERVICE -> DCLSRVTMPL -> PROTOCOLSERVICE."""
        return self.ref(self.ref(service, "tmplref"), "tmplref")

    def request_bytes(self, diaginst, service):
        """Build the constant leading bytes of the request (SID, sub-function, identifier)."""
        proto = self.protocol_service(service)
        req = proto.find("REQ") if proto is not None else None
        if req is None:
            return None, ""
        static_values = {}
        for sv in diaginst.findall("STATICVALUE"):
            shstatic = self.ref(sv, "shstaticref")
            for comp_ref in shstatic.findall("STATICCOMPREF") if shstatic is not None else []:
                static_values[comp_ref.get("idref")] = int(sv.get("v"))
        sid = None
        parts = []
        for comp in req:
            if comp.tag == "CONSTCOMP":
                n_bytes = max(1, int(comp.get("bl", 8)) // 8)
                value = int(comp.get("v"))
                if comp.get("spec") == "sid" and sid is None:
                    sid = value
                parts.append(to_hex(value, n_bytes))
            elif comp.tag == "STATICCOMP":
                bits, _ = self.bit_length(comp.get("dtref"))
                n_bytes = max(1, (bits or 8) // 8)
                value = static_values.get(comp.get("id"))
                parts.append(to_hex(value, n_bytes) if value is not None else "XX" * n_bytes)
            elif comp.tag in ("NAME", "QUAL", "DESC"):
                continue
            else:
                break  # first variable data component ends the constant prefix
        return sid, " ".join(parts)

    # ---------- structure walking ----------

    def variants(self):
        return self.doc.find("ECU").findall("VAR")

    def diag_instances(self, variant):
        """Yield (diag class name, DIAGINST) for instances defined in or referenced by a variant."""
        for child in variant:
            if child.tag == "DIAGCLASS":
                for inst in child.findall("DIAGINST"):
                    yield text(child), inst
            elif child.tag == "DIAGINST":
                yield text(self.ref(child, "tmplref")), child
            elif child.tag == "DIAGINSTREF":
                inst = self.ref(child, "idref")
                if inst is not None:
                    yield text(self.ref(inst, "tmplref")), inst

    def parameters(self, diaginst):
        """Return (group name, [element names]) for each data block of an instance."""
        out = []
        for cont in diaginst.findall("SIMPLECOMPCONT"):
            names = [text(e) for e in cont.iter("DATAOBJ")]
            if names:
                out.append((text(self.ref(cont, "shproxyref")), names))
        return out

    # ---------- XML builders ----------

    def services(self):
        section = ET.Element("Services")
        for variant in self.variants():
            group = add(section, "Variant", name=text(variant))
            for diag_class, inst in self.diag_instances(variant):
                for service in inst.findall("SERVICE"):
                    sid, request = self.request_bytes(inst, service)
                    sessions, security = self.access(service)
                    addressing = [a for a, flag in (("Physical", "phys"), ("Functional", "func")) if service.get(flag) == "1"]
                    add(group, "Service",
                        sid=to_hex(sid, 1) if sid is not None else "",
                        request=request,
                        diagClass=diag_class,
                        instance=text(inst),
                        name=text(service),
                        qualifier=service.findtext("SHORTCUTQUAL") or service.findtext("QUAL") or "",
                        protocolService=text(self.protocol_service(service)),
                        addressing=", ".join(addressing),
                        sessions=sessions,
                        security=security)
        return section

    def dids(self):
        # Which variants list each DID, and which services read or write it.
        did_variants = {}
        for variant in self.variants():
            for ref in variant.iter("DIDREF"):
                did_variants.setdefault(ref.get("didRef"), []).append(text(variant))
        access = {}  # (variant, DID number, SID) -> (sessions, security)
        for variant in self.variants():
            for _, inst in self.diag_instances(variant):
                sv = inst.find("STATICVALUE")
                if sv is None:
                    continue
                for service in inst.findall("SERVICE"):
                    sid, _ = self.request_bytes(inst, service)
                    if sid in (READ_DID_SID, WRITE_DID_SID):
                        access.setdefault((text(variant), int(sv.get("v")), sid), self.access(service))

        section = ET.Element("DIDs")
        for did in self.doc.find("DIDS").findall("DID"):
            number = int(did.get("n"))
            variants = did_variants.get(did.get("id"), [])
            node = add(section, "DID", id=to_hex(number, 2), name=text(did), qualifier=did.findtext("QUAL") or "",
                       variants=", ".join(variants))
            for label, sid in (("Read", READ_DID_SID), ("Write", WRITE_DID_SID)):
                found = next((access[k] for k in ((v, number, sid) for v in variants) if k in access), None)
                if found:
                    add(node, label, sid=to_hex(sid, 1), sessions=found[0], security=found[1])
            min_bits, max_bits = self.add_data_elements(node, did.find("STRUCTURE"))
            if min_bits is None:
                length = "?"
            elif min_bits == max_bits:
                length = f"{min_bits // 8}" if min_bits % 8 == 0 else f"{min_bits} bit"
            else:
                length = f"{min_bits // 8}..{max_bits // 8}"
            node.set("length", length)
        return section

    def add_data_elements(self, parent, structure, start_bit=0):
        """Add a DataElement per element of a structure, descending into nested structures.

        Returns the (min, max) bit length of the structure, or (None, None) if unknown."""
        position, min_total, max_total = start_bit, 0, 0
        for obj in structure if structure is not None else []:
            if obj.tag not in ("DATAOBJ", "STRUCT", "GAPDATAOBJ"):
                continue
            if obj.tag == "GAPDATAOBJ":
                lo = hi = int(obj.get("bl", 0))
                dtype = None
            else:
                lo, hi = self.bit_length(obj.get("dtref"))
                dtype = self.ids.get(obj.get("dtref"))
            element = add(parent, "DataElement",
                          name=text(obj),
                          qualifier=obj.findtext("QUAL") or "",
                          startByte=position // 8 if position is not None else "?",
                          startBit=position % 8 if position is not None else "?",
                          bitLength="?" if lo is None else (lo if lo == hi else f"{lo}..{hi}"))
            if obj.tag == "GAPDATAOBJ":
                element.set("dataType", "Gap (reserved)")
            else:
                self.add_type_details(element, dtype)
            # Children of an inline STRUCT or of a structured data type, positioned inside the parent.
            children = obj if obj.tag == "STRUCT" else dtype
            if children is not None and children.tag in ("STRUCT", "STRUCTDT", "EOSITERDT", "STATICITERRDT"):
                inner = children.find("STRUCT") if children.tag == "STRUCTDT" else children
                self.add_data_elements(element, inner, position)
            if lo is None or min_total is None:
                min_total = max_total = None
            else:
                min_total, max_total = min_total + lo, max_total + hi
            position = position + lo if position is not None and lo is not None and lo == hi else None
        return min_total, max_total

    def add_type_details(self, element, dtype):
        """Describe a CANdela data type: kind, coded and physical value, conversion, unit and value table."""
        if dtype is None:
            return
        cvt, pvt = dtype.find("CVALUETYPE"), dtype.find("PVALUETYPE")
        set_attrs(element,
                  dataType=f"{DATA_TYPE_KINDS.get(dtype.tag, dtype.tag)}: {text(dtype)}",
                  codedType=self.value_type(cvt),
                  byteOrder=BYTE_ORDERS.get(cvt.get("bo"), cvt.get("bo")) if cvt is not None else "",
                  physicalType=self.value_type(pvt),
                  displayFormat=DISPLAY_FORMATS.get(pvt.get("df"), pvt.get("df")) if pvt is not None else "",
                  unit=(pvt.findtext("UNIT") or "") if pvt is not None else "")
        for comp in dtype.findall("COMP"):
            raw = f"{comp.get('s', '')}..{comp.get('e', '')}" if comp.get("s") or comp.get("e") else ""
            add(element, "Conversion", raw=raw, factor=comp.get("f", "1"), divisor=comp.get("div", ""),
                offset=comp.get("o", "0"))
        for mapping in dtype.findall("TEXTMAP"):
            add(element, "Value", raw=value_range(mapping)).text = text(mapping, "TEXT")
        for excluded in dtype.findall("EXCL"):
            add(element, "Invalid", raw=value_range(excluded)).text = excluded.get("inv", "invalid")

    @staticmethod
    def value_type(vt):
        """Describe a CVALUETYPE/PVALUETYPE, e.g. 'Unsigned 16 bit' or 'ASCII 8 bit x 17'."""
        if vt is None:
            return ""
        desc = f"{ENCODINGS.get(vt.get('enc'), vt.get('enc'))} {vt.get('bl')} bit"
        if vt.get("qty") == "field":
            lo, hi = vt.get("minsz"), vt.get("maxsz")
            desc += f" x {lo}" if lo == hi else f" x {lo}..{hi}"
        return desc

    def dtcs(self):
        section = ET.Element("DTCs")
        for record_dt in self.root.iter("RECORDDT"):
            records = record_dt.findall("RECORD")
            if not records:
                continue
            variant, memory = "Record pool (no variant)", ""
            node = self.parent.get(record_dt)
            while node is not None:
                if node.tag == "DIAGINST":
                    memory = text(node)
                elif node.tag == "VAR":
                    variant = text(node)
                    break
                node = self.parent.get(node)
            group = add(section, "FaultMemory", variant=variant, name=memory)
            for record in records:
                number = int(record.get("v"))
                dtc = add(group, "DTC", id=to_hex(number, 3), code=dtc_code(number), description=text(record, "TEXT"))
                for item in record:
                    tmpl = self.ref(item, "idref")
                    if tmpl is None:
                        continue
                    if item.tag == "TRRECORDITEM":
                        value = text(item, "TEXT")
                    elif item.tag == "ENUMRECORDITEM":
                        labels = {e.get("v"): text(e, None) for e in tmpl.findall("ETAG")}
                        value = labels.get(item.get("v"), item.get("v"))
                    else:
                        value = item.get("v")
                    if value:
                        add(dtc, "Field", name=text(tmpl)).text = value
        return section

    def rids(self):
        section = ET.Element("RIDs")
        for variant in self.variants():
            group = None
            for diag_class, inst in self.diag_instances(variant):
                services = inst.findall("SERVICE")
                sv = inst.find("STATICVALUE")
                if not services or sv is None:
                    continue
                requests = [(s, *self.request_bytes(inst, s)) for s in services]
                if not any(sid == ROUTINE_CONTROL_SID for _, sid, _ in requests):
                    continue
                if group is None:
                    group = add(section, "Variant", name=text(variant))
                sessions, security = self.access(services[0])
                rid = add(group, "RID", id=to_hex(int(sv.get("v")), 2), name=text(inst),
                          qualifier=inst.findtext("QUAL") or "", diagClass=diag_class,
                          sessions=sessions, security=security)
                if text(inst, "DESC"):
                    add(rid, "Description").text = text(inst, "DESC")
                for service, _, request in requests:
                    add(rid, "Request", subFunction=text(service)).text = request
                for group_name, names in self.parameters(inst):
                    params = add(rid, "Parameters", group=group_name)
                    for name in names:
                        add(params, "Parameter", name=name)
        return section


def set_attrs(elem, **attrs):
    """Set attributes, leaving out empty ones so the XML stays short."""
    for key, value in attrs.items():
        if value not in (None, ""):
            elem.set(key, str(value))


def add(parent, tag, **attrs):
    child = ET.SubElement(parent, tag)
    set_attrs(child, **attrs)
    return child


def build_export(cdd_path):
    """Parse a CDD file into the export tree. Returns (root element, counts per section)."""
    cdd = CddFile(cdd_path)
    root = ET.Element("CddExport", source=cdd.path.name, ecu=text(cdd.doc.find("ECU")))
    sections = {"Services": cdd.services(), "DIDs": cdd.dids(), "DTCs": cdd.dtcs(), "RIDs": cdd.rids()}
    counts = {
        "Services": len(sections["Services"].findall("Variant/Service")),
        "DIDs": len(sections["DIDs"].findall("DID")),
        "DTCs": len(sections["DTCs"].findall("FaultMemory/DTC")),
        "RIDs": len(sections["RIDs"].findall("Variant/RID")),
    }
    for name, section in sections.items():
        section.set("count", str(counts[name]))
        root.append(section)
    return root, counts


def write_xml(root, xml_path):
    ET.indent(root, space="  ")
    ET.ElementTree(root).write(xml_path, encoding="utf-8", xml_declaration=True)


# ---------- Excel output ----------

def headers(attrs, names):
    """Rename XML attributes to column titles, e.g. {'sid': ...} -> {'SID': ...}."""
    return {names.get(key, key[:1].upper() + key[1:]): value for key, value in attrs.items()}


SERVICE_COLUMNS = {"sid": "SID", "request": "Request", "diagClass": "Diag Class", "instance": "Diag Instance",
                   "name": "Service", "protocolService": "Protocol Service"}
DID_COLUMNS = {"id": "DID", "length": "Length (bytes)"}
ELEMENT_COLUMNS = {"name": "Data Element", "qualifier": "Element Qualifier", "startByte": "Start Byte",
                   "startBit": "Start Bit", "bitLength": "Bit Length", "dataType": "Data Type",
                   "codedType": "Coded Type", "byteOrder": "Byte Order", "physicalType": "Physical Type",
                   "displayFormat": "Display Format"}
RID_COLUMNS = {"id": "RID", "diagClass": "Diag Class"}


def element_rows(parent, prefix=""):
    """Flatten (nested) DataElements into rows, naming children 'Parent / Child'."""
    for element in parent.findall("DataElement"):
        row = headers(element.attrib, ELEMENT_COLUMNS)
        row["Data Element"] = prefix + element.get("name", "")
        row["Conversion"] = "\n".join(
            (f"raw {c.get('raw')}: " if c.get("raw") else "")
            + ", ".join(f"{k} {c.get(k)}" for k in ("factor", "divisor", "offset") if c.get(k))
            for c in element.findall("Conversion"))
        row["Unit"] = row.pop("Unit", "")
        row["Value Table"] = "\n".join(f"{v.get('raw')} = {v.text}" for v in element.findall("Value"))
        row["Invalid Values"] = "\n".join(f"{v.get('raw')} ({v.text})" for v in element.findall("Invalid"))
        yield row
        yield from element_rows(element, row["Data Element"] + " / ")


def excel_sheets(root):
    """Turn the export tree into one list of row dicts per sheet."""
    services = [{"Variant": variant.get("name"), **headers(service.attrib, SERVICE_COLUMNS)}
                for variant in root.find("Services") for service in variant]

    dids = []
    for did in root.find("DIDs"):
        did_columns = headers({k: v for k, v in did.attrib.items()}, DID_COLUMNS)
        access = {}
        for label, sid in (("Read", "22"), ("Write", "2E")):
            node = did.find(label)
            access[f"{label} ({sid})"] = "Yes" if node is not None else "No"
            access[f"{label} Sessions"] = node.get("sessions", "") if node is not None else ""
            access[f"{label} Security"] = node.get("security", "") if node is not None else ""
        for element in list(element_rows(did)) or [{}]:
            dids.append({**did_columns, **element, **access})

    dtcs = []
    for memory in root.find("DTCs"):
        for dtc in memory:
            row = {"Variant": memory.get("variant"), "Fault Memory": memory.get("name", ""),
                   "DTC": dtc.get("id"), "DTC Code": dtc.get("code"), "Description": dtc.get("description", "")}
            row.update((field.get("name"), field.text) for field in dtc.findall("Field"))
            dtcs.append(row)

    rids = []
    for variant in root.find("RIDs"):
        for rid in variant:
            row = {"Variant": variant.get("name"), **headers(rid.attrib, RID_COLUMNS)}
            row["Requests"] = "\n".join(f"{r.get('subFunction')}: {r.text}" for r in rid.findall("Request"))
            row["Parameters"] = "\n".join(
                f"{p.get('group')}: {', '.join(x.get('name') for x in p)}" for p in rid.findall("Parameters"))
            row["Description"] = rid.findtext("Description", "")
            rids.append(row)

    return {"Service": services, "DID": dids, "DTC": dtcs, "RID": rids}


NUMERIC_COLUMNS = {"Start Byte", "Start Bit", "Bit Length", "Length (bytes)"}


def excel_value(sheet, column, value):
    """Store counts and lengths as numbers; hex identifiers such as DID, SID or DTC stay text."""
    numeric = column in NUMERIC_COLUMNS or (sheet == "DTC" and column not in ("DTC", "DTC Code"))
    if numeric and isinstance(value, str) and value.isdigit():
        return int(value)
    return value if value is not None else ""


def write_excel(root, xlsx_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Excel output needs openpyxl: pip install openpyxl")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in excel_sheets(root).items():
        sheet = workbook.create_sheet(title)
        # The fullest row gives the column order; columns only some rows have go after it.
        columns = []
        for row in sorted(rows, key=len, reverse=True):
            columns.extend(c for c in row if c not in columns)
        sheet.append(columns)
        for row in rows:
            sheet.append([excel_value(title, c, row.get(c)) for c in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(columns, start=1):
            longest = max([len(column)] + [max((len(line) for line in str(r.get(column) or "").split("\n")), default=0)
                                           for r in rows])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
    workbook.save(xlsx_path)


WRITERS = {"xml": (write_xml, ".xml"), "excel": (write_excel, ".xlsx")}


def export(cdd_path, output_path, output_format="xml"):
    root, counts = build_export(cdd_path)
    WRITERS[output_format][0](root, output_path)
    return counts


def main(argv=None):
    parser = argparse.ArgumentParser(description="Export DIDs, DTCs, RIDs and services of a CDD file to XML or Excel.")
    parser.add_argument("cdd", help="path to the .cdd file")
    parser.add_argument("-f", "--format", choices=sorted(WRITERS), default="xml",
                        help="output format (default: xml)")
    parser.add_argument("-o", "--output", help="output file path (default: <cdd name>.xml or .xlsx next to the CDD)")
    args = parser.parse_args(argv)
    output = args.output or str(Path(args.cdd).with_suffix(WRITERS[args.format][1]))
    counts = export(args.cdd, output, args.format)
    print(f"Wrote {output}")
    for name, count in counts.items():
        print(f"  {name}: {count}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
